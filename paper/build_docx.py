from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
import re
D=Path(__file__).parent
text=(D/'manuscript.txt').read_text()
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.73);sec.bottom_margin=Inches(.70)
sec.left_margin=Inches(.76);sec.right_margin=Inches(.76)
sec.header_distance=Inches(.30);sec.footer_distance=Inches(.30)

def set_font(st,latin='Times New Roman',size=11,bold=False,east='Noto Serif CJK SC'):
 st.font.name=latin;st.font.size=Pt(size);st.font.bold=bold;st.font.color.rgb=RGBColor(0,0,0)
 pr=st.element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
 if rf is None: rf=OxmlElement('w:rFonts');pr.insert(0,rf)
 for k,v in [('ascii',latin),('hAnsi',latin),('eastAsia',east)]:rf.set(qn('w:'+k),v)
 if rf.get(qn('w:asciiTheme')): del rf.attrib[qn('w:asciiTheme')]
for name in ['Normal','Body Text']:
 set_font(doc.styles[name],size=10.7)
 pf=doc.styles[name].paragraph_format;pf.line_spacing=Pt(16.3);pf.space_after=Pt(5);pf.widow_control=True
set_font(doc.styles['Title'],size=24,bold=True,east='Noto Sans CJK SC')
doc.styles['Title'].paragraph_format.space_after=Pt(7)
doc.styles['Title'].paragraph_format.line_spacing=Pt(28)
set_font(doc.styles['Subtitle'],size=13,bold=False,east='Noto Sans CJK SC')
doc.styles['Subtitle'].paragraph_format.space_after=Pt(7)
doc.styles['Subtitle'].font.italic=False
doc.styles['Subtitle'].paragraph_format.line_spacing=Pt(19)
for name,size in [('Heading 1',15),('Heading 2',11.7)]:
 set_font(doc.styles[name],size=size,bold=True,east='Noto Sans CJK SC')
 pf=doc.styles[name].paragraph_format;pf.space_before=Pt(12 if name=='Heading 1' else 8);pf.space_after=Pt(5);pf.keep_with_next=True;pf.keep_together=True;pf.line_spacing=Pt(20 if name=='Heading 1' else 17)
set_font(doc.styles['Caption'],size=9.1)
doc.styles['Caption'].paragraph_format.line_spacing=Pt(12);doc.styles['Caption'].paragraph_format.space_after=Pt(8)
for name,size in [('Equation',10),('Reference',9.0),('Step',10.2),('Metadata',9.5)]:
 if name not in doc.styles:doc.styles.add_style(name,1)
 set_font(doc.styles[name],size=size)
 pf=doc.styles[name].paragraph_format;pf.space_after=Pt(5);pf.widow_control=True;pf.line_spacing=Pt(11.3 if name=='Reference' else 15)
if 'Equation' in doc.styles:
 doc.styles['Equation'].paragraph_format.left_indent=Inches(.18)
 doc.styles['Equation'].paragraph_format.space_before=Pt(3)
 doc.styles['Equation'].paragraph_format.space_after=Pt(7)
 doc.styles['Equation'].paragraph_format.keep_together=True
for st in doc.styles:
 for border in list(st.element.xpath('.//w:pBdr')):border.getparent().remove(border)
# footer page field, useful for a multi-page manuscript
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=p.add_run();r.font.name='Times New Roman';r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');r._r.addnext(fld)
# readable table helpers

def shade(cell,color):
 pr=cell._tc.get_or_add_tcPr();x=OxmlElement('w:shd');x.set(qn('w:fill'),color);pr.append(x)
def table(headers,rows,widths):
 t=doc.add_table(rows=1, cols=len(headers));t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for j,h in enumerate(headers):t.rows[0].cells[j].text=h
 for rr in rows:
  cs=t.add_row().cells
  for j,v in enumerate(rr):cs[j].text=v
 for i,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if i==0:
   rpt=OxmlElement('w:tblHeader');trPr.append(rpt)
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=cell._tc.get_or_add_tcPr()
   margins=OxmlElement('w:tcMar')
   for side,val in [('top',55),('left',90),('bottom',55),('right',90)]:
    x=OxmlElement('w:'+side);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');margins.append(x)
   tcpr.append(margins)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
   tcpr.append(borders)
   shade(cell,'DFE8F2' if i==0 else ('F6F8FA' if i%2==0 else 'FFFFFF'))
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=Pt(13);p.paragraph_format.keep_with_next=False
    if j==0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:
     r.font.name='Times New Roman';r.font.size=Pt(9.2);r.font.bold=(i==0);r.font.color.rgb=RGBColor(0,0,0)
     r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
 return t
arms=[
 ['B_native','a₀','无回执 无审阅包','原生消息基线'],
 ['B','a₀','仅真实控制器回执','测量回执装置效应'],
 ['S','自我重考虑','按预定自我修订协议','额外思考对照；真实生成待实现'],
 ['A','a₁','只有事实回执','不呈现原始审阅包'],
 ['P₀','a₁','每次保留原始包','对应冻结适配器的脚本 P'],
 ['P','a₁','每次保留包和中性头','主比较基线；包装版本待实现'],
 ['C','a₁','首次同 P；随后按公开条件变更状态头','主处理；原包文本保留'],
 ['R','a₁','首次同 P；随后不再附带整个包','一次呈现基线；可丢失有用信息'],
]
readiness=[
 ['后端恢复','7 个脚本原运行见证；14 次独立进程恢复；42 个检查点工具对和 84 次 READ 转移','固定提交上的窄工具及状态恢复；非自然轨迹或效果样本'],
 ['内部干预','36 项测试；B/S/A/P 的内部事件和逻辑记账','S 只验证预写的格式错误回退；没有实际自我修订'],
 ['回执适配','37 项测试；6 次独立进程重启检查','检查点 0 与两次脚本执行者边界；原包 P₀ 的呈现契约'],
 ['本稿主方法','新 P、C、R、共同首段后的新检查点、真实续写','尚未实现或验证；不能从现有通过数推定可用'],
]
stepn=0;inrefs=False
for block in text.split('\n\n'):
 block=block.strip()
 if not block:continue
 # headings can be adjacent to paragraph in same block
 lines=block.splitlines()
 while lines and lines[0].startswith(('TITLE:','SUBTITLE:','META:')):
  l=lines.pop(0)
  if l.startswith('TITLE:'):doc.add_paragraph(l[6:].strip(),'Title')
  elif l.startswith('SUBTITLE:'):doc.add_paragraph(l[9:].strip(),'Subtitle')
  else:doc.add_paragraph(l[5:].strip(),'Metadata')
 if not lines:continue
 if lines[0].startswith('# '):
  head=lines.pop(0)[2:];doc.add_paragraph(head,'Heading 1');inrefs=(head=='参考文献')
  stepn=0
 if lines and lines[0].startswith('## '):
  head=lines.pop(0)[3:];doc.add_paragraph(head,'Heading 2');stepn=0
 if not lines:continue
 block='\n'.join(lines)
 if block=='TABLE_ARMS':
  doc.add_paragraph('表 1 预定比较分支与当前实现的关系','Caption').paragraph_format.keep_with_next=True
  table(['分支','执行动作','建议呈现','目的与状态'],arms,[.76,.78,2.76,2.68]);doc.add_paragraph().paragraph_format.space_after=Pt(0);continue
 if block=='TABLE_READINESS':
  doc.add_paragraph('表 2 软件验证与科学证据的边界','Caption').paragraph_format.keep_with_next=True
  table(['层级','已有验证','能支持的结论或限制'],readiness,[1.0,2.64,3.34]);doc.add_paragraph().paragraph_format.space_after=Pt(0);continue
 if block.startswith('FIGURE:'):
  p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(3)
  r=p.add_run();pic=r.add_picture(str(D/'reviewer_advice_overview.png'),width=Inches(6.97))
  pic._inline.docPr.set('descr','Proposed experiment with shared reviewer action, persistent and consumed-status packet projections, and an evaluator-only goal scoring boundary. No empirical outcomes shown.')
  if '\nCAPTION:' in block:doc.add_paragraph(block.split('\nCAPTION:',1)[1].strip(),'Caption')
  continue
 if block.startswith('EQUATION:'):
  for l in block.splitlines():
   if l.startswith('EQUATION:'):doc.add_paragraph(l[len('EQUATION:'):].strip(),'Equation')
  continue
 if block.startswith('STEP:'):
  for l in block.splitlines():
   if l.startswith('STEP:'):
    stepn+=1;doc.add_paragraph(f'{stepn}. '+l[len('STEP:'):].strip(),'Step')
  continue
 p=doc.add_paragraph(block,'Reference' if inrefs else 'Normal')
 if inrefs:p.paragraph_format.keep_together=True
# Editable typographic subscripts for the manuscript notation.
pattern=re.compile(r"(ΔY|Δ|φ|s|h|o|S|L|U|Y|V|D|E|B)_(Dref,q,ω|root|pub|ref|native|P₀A|PA|PC|PR|AB|t|u|T|x|g)")
for pp in doc.paragraphs:
 if pp.style.name in ['Title','Subtitle','Heading 1','Heading 2','Reference','Metadata','Caption']:
  continue
 original=pp.text
 if not pattern.search(original):continue
 for child in list(pp._p):
  if child.tag != qn('w:pPr'):pp._p.remove(child)
 pos=0
 for m in pattern.finditer(original):
  pp.add_run(original[pos:m.start()]);pp.add_run(m.group(1))
  rr=pp.add_run(m.group(2));rr.font.subscript=True
  pos=m.end()
 pp.add_run(original[pos:])
# metadata
props=doc.core_properties;props.title='Reviewer Advice Lifetimes and Goal Preservation in Multi Turn Agents';props.subject='Pre-experiment research manuscript on reviewer advice scope and achieved-goal preservation';props.author='jiuxiyan';props.keywords='reviewer feedback; tool agents; goal preservation; controlled experiment';props.comments=''
# enforce black title and no paragraph borders
for p in doc.paragraphs:
 if p.style.name in ['Title','Subtitle','Heading 1','Heading 2']:
  for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
  pp=p._p.find(qn('w:pPr'))
  if pp is not None:
   for n in list(pp):
    if n.tag==qn('w:pBdr'):pp.remove(n)
out=D/'reviewer_advice_lifetimes_preexperiment.docx';doc.save(out)
print(out)
print('paragraphs',len(doc.paragraphs),'tables',len(doc.tables))
